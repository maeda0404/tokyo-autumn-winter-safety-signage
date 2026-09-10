#!/usr/bin/env python3
"""
京王笹塚 秋冬安全サイネージ用 status.json 生成スクリプト

取得情報
- 気象庁 VPWS50（集約定時通報）: 東京地方の強風注意報・暴風警報・乾燥注意報
- 気象庁 府県天気予報: 今夜～明朝の最低気温
- 東京都感染症情報センター: インフルエンザ定点当たり患者報告数

VPWS50を警報・注意報の主データとし、取得成功時はその結果を確定する。
旧 warning/130000.json は使用しない。
"""

import gzip
import io
import json
import re
import time
import urllib.request
import xml.etree.ElementTree as ET
from datetime import date, datetime, timedelta, timezone
from html import unescape
from pathlib import Path
from urllib.parse import urljoin

from openpyxl import load_workbook

JST = timezone(timedelta(hours=9))
ROOT_DIR = Path(__file__).resolve().parents[1]
STATUS_FILE = ROOT_DIR / "data" / "status.json"

REGULAR_FEED_URL = "https://www.data.jma.go.jp/developer/xml/feed/regular.xml"
FORECAST_URL = "https://www.jma.go.jp/bosai/forecast/data/forecast/130000.json"
FLU_TOP_URL = "https://idsc.tmiph.metro.tokyo.lg.jp/diseases/flu/flu/"

# VPWS50は全国集約電文。東京地方を対象とする。
VPWS50_TYPE = "VPWS50"
VPWS50_AREA = "010000"
TARGET_AREA_CODES = {"130010"}
TARGET_AREA_NAMES = {"東京地方"}

# 東京地点の気温コード
TOKYO_POINT_CODE = "44132"

DATA_URL_RE = re.compile(
    r"https://www\.data\.jma\.go\.jp/developer/xml/data/"
    r"(\d{14})_\d+_([A-Z0-9]+)_(\d{6})\.xml"
)

INACTIVE_STATUSES = {
    "解除",
    "なし",
    "発表警報・注意報はなし",
    "警報・注意報はなし",
}


def fetch_bytes(url, retries=3):
    last_error = None
    for attempt in range(1, retries + 1):
        try:
            separator = "&" if "?" in url else "?"
            cache_buster = int(datetime.now(timezone.utc).timestamp())
            request = urllib.request.Request(
                f"{url}{separator}_={cache_buster}",
                headers={
                    "User-Agent": "Mozilla/5.0 Tokyo-Autumn-Winter-Safety-Signage/3.0",
                    "Accept": "application/xml, application/json, text/html, */*",
                    "Cache-Control": "no-cache",
                    "Pragma": "no-cache",
                },
            )
            with urllib.request.urlopen(request, timeout=30) as response:
                raw = response.read()
            return gzip.decompress(raw) if raw[:2] == b"\x1f\x8b" else raw
        except Exception as error:
            last_error = error
            print(f"通信失敗 {attempt}/{retries}: {url}: {error}")
            if attempt < retries:
                time.sleep(5 * attempt)
    raise last_error


def fetch_text(url):
    return fetch_bytes(url).decode("utf-8", errors="replace")


def fetch_json(url):
    return json.loads(fetch_bytes(url).decode("utf-8"))


def html_to_text(html):
    text = re.sub(r"<script[\s\S]*?</script>", " ", html, flags=re.IGNORECASE)
    text = re.sub(r"<style[\s\S]*?</style>", " ", html, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", unescape(text)).strip()


def load_previous_status():
    if not STATUS_FILE.exists():
        return {"influenza": {}, "weather": {}}
    try:
        value = json.loads(STATUS_FILE.read_text(encoding="utf-8"))
        return value if isinstance(value, dict) else {"influenza": {}, "weather": {}}
    except Exception as error:
        print("前回データ読込エラー:", error)
        return {"influenza": {}, "weather": {}}


# ---------------------------------------------------------------------------
# VPWS50
# ---------------------------------------------------------------------------
def local_name(element):
    return element.tag.rsplit("}", 1)[-1]


def direct_text(element, name):
    for child in element:
        if local_name(child) == name:
            return (child.text or "").strip()
    return ""


def select_latest_vpws50_url(feed_bytes):
    text = feed_bytes.decode("utf-8", errors="replace")
    candidates = []
    for timestamp, message_type, area in DATA_URL_RE.findall(text):
        if message_type == VPWS50_TYPE and area == VPWS50_AREA:
            candidates.append(
                (
                    timestamp,
                    "https://www.data.jma.go.jp/developer/xml/data/"
                    f"{timestamp}_0_{message_type}_{area}.xml",
                )
            )
    return max(candidates, key=lambda item: item[0]) if candidates else None


def item_matches_tokyo(item):
    for element in item.iter():
        if local_name(element) != "Area":
            continue
        area_name = direct_text(element, "Name")
        area_code = direct_text(element, "Code")
        if area_code in TARGET_AREA_CODES or area_name in TARGET_AREA_NAMES:
            return True
    return False


def parse_vpws50(xml_bytes):
    """東京地方の強風・暴風・乾燥を解析する。VPWS50の空Statusは有効。"""
    root = ET.fromstring(xml_bytes)
    report_datetime = ""
    for element in root.iter():
        if local_name(element) == "ReportDateTime" and element.text:
            report_datetime = element.text.strip()
            break

    active_names = []
    wind_active = False
    dry_active = False

    for item in (element for element in root.iter() if local_name(element) == "Item"):
        if not item_matches_tokyo(item):
            continue
        for kind in (element for element in item.iter() if local_name(element) == "Kind"):
            name = direct_text(kind, "Name")
            status = direct_text(kind, "Status")

            # VPWS50は現在状況のスナップショットなので空Statusも有効。
            if not name or status in INACTIVE_STATUSES:
                continue

            if any(word in name for word in ("強風注意報", "暴風警報", "暴風特別警報", "暴風雪警報", "暴風雪特別警報", "風雪注意報")):
                wind_active = True
                active_names.append(name)

            if "乾燥注意報" in name:
                dry_active = True
                active_names.append(name)

    return {
        "windActive": wind_active,
        "dryActive": dry_active,
        "activeNames": list(dict.fromkeys(active_names)),
        "reportDateTime": report_datetime,
    }


def get_vpws50_weather_alerts():
    selected = select_latest_vpws50_url(fetch_bytes(REGULAR_FEED_URL))
    if not selected:
        raise ValueError("regular.xmlからVPWS50を取得できません")
    _, xml_url = selected
    result = parse_vpws50(fetch_bytes(xml_url))
    result["sourceFile"] = xml_url.rsplit("/", 1)[-1]
    result["sourceType"] = VPWS50_TYPE
    result["emptyStatusAccepted"] = True
    result["resultLocked"] = True
    return result


# ---------------------------------------------------------------------------
# 最低気温
# ---------------------------------------------------------------------------
def parse_jma_datetime(value):
    return datetime.fromisoformat(value).astimezone(JST)


def get_overnight_minimum_temperature():
    """翌日0時の東京地点気温を優先し、なければ翌日の週間最低気温を使う。"""
    data = fetch_json(FORECAST_URL)
    target_date = datetime.now(JST).date() + timedelta(days=1)

    for report in data:
        for series in report.get("timeSeries", []):
            times = series.get("timeDefines", [])
            for area_data in series.get("areas", []):
                if str(area_data.get("area", {}).get("code", "")) != TOKYO_POINT_CODE:
                    continue
                for time_text, value in zip(times, area_data.get("temps", [])):
                    try:
                        forecast_time = parse_jma_datetime(time_text)
                        temperature = float(value)
                    except (TypeError, ValueError):
                        continue
                    if forecast_time.date() == target_date and forecast_time.hour == 0:
                        return temperature, target_date, "shortForecast00"

    for report in data:
        for series in report.get("timeSeries", []):
            times = series.get("timeDefines", [])
            for area_data in series.get("areas", []):
                area = area_data.get("area", {})
                code = str(area.get("code", ""))
                name = str(area.get("name", ""))
                if code not in {TOKYO_POINT_CODE, "130010"} and "東京" not in name:
                    continue
                for time_text, value in zip(times, area_data.get("tempsMin", [])):
                    try:
                        forecast_time = parse_jma_datetime(time_text)
                        temperature = float(value)
                    except (TypeError, ValueError):
                        continue
                    if forecast_time.date() == target_date:
                        return temperature, target_date, "weeklyTempsMin"

    return None, target_date, "notFound"


# ---------------------------------------------------------------------------
# インフルエンザ
# ---------------------------------------------------------------------------
def determine_flu_level(value):
    if value is None:
        return "確認中"
    if value >= 30:
        return "警報レベル"
    if value >= 10:
        return "注意報レベル"
    if value >= 1:
        return "流行中"
    return "非流行"


def find_flu_excel_url(top_html):
    links = re.findall(
        r"<a\b[^>]*href=[\"']([^\"']+)[\"'][^>]*>([\s\S]*?)</a>",
        top_html,
        flags=re.IGNORECASE,
    )
    for href, label_html in links:
        url = urljoin(FLU_TOP_URL, unescape(href))
        label = html_to_text(label_html)
        lower = url.lower()
        if (".xlsx" in lower or ".xls" in lower) and (
            "hasseidoko" in lower or "患者報告数" in label
        ):
            return url
    raise ValueError("東京都公式ページから患者報告数ExcelのURLを取得できません")


def normalize(value):
    return "" if value is None else re.sub(r"\s+", "", str(value)).strip()


def number(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = normalize(value).replace(",", "").replace("人", "")
    return float(text) if re.fullmatch(r"-?\d+(?:\.\d+)?", text) else None


def parse_week_header(value):
    match = re.fullmatch(r"(?:第)?(\d{1,2})(?:w|週)", normalize(value).lower())
    if not match:
        return None
    week = int(match.group(1))
    return week if 1 <= week <= 53 else None


def find_latest_record(workbook):
    now = datetime.now(JST)
    candidates = []

    for sheet in workbook.worksheets:
        header_row = None
        week_columns = {}
        for row in range(1, min(sheet.max_row, 30) + 1):
            found = {}
            for column in range(1, sheet.max_column + 1):
                week = parse_week_header(sheet.cell(row=row, column=column).value)
                if week is not None:
                    found[column] = week
            if len(found) > len(week_columns):
                header_row, week_columns = row, found
        if not header_row or not week_columns:
            continue

        total_row = None
        sentinel_row = None
        for row in range(header_row + 1, sheet.max_row + 1):
            row_labels = [
                normalize(sheet.cell(row=row, column=column).value).lower()
                for column in range(1, min(sheet.max_column, 8) + 1)
            ]
            if any(label in {"total", "合計"} for label in row_labels):
                total_row = row
            if any("定点数" in label or "テイテンスウ" in label for label in row_labels):
                sentinel_row = row

        if total_row is None or sentinel_row is None:
            continue

        sheet_text = " ".join(
            normalize(sheet.cell(row=row, column=column).value)
            for row in range(1, min(sheet.max_row, 8) + 1)
            for column in range(1, min(sheet.max_column, 12) + 1)
        )
        season_match = re.search(r"(20\d{2})[-－](\d{2})", sheet_text)
        season_start_year = int(season_match.group(1)) if season_match else now.year - 1

        for column, week in week_columns.items():
            total = number(sheet.cell(row=total_row, column=column).value)
            points = number(sheet.cell(row=sentinel_row, column=column).value)
            if total is None or points is None or points <= 0:
                continue
            year = season_start_year if week >= 36 else season_start_year + 1
            try:
                week_end = date.fromisocalendar(year, week, 7)
            except ValueError:
                continue
            if week_end > now.date():
                continue
            candidates.append({
                "year": year,
                "week": week,
                "value": round(total / points, 2),
                "reportedCases": int(total),
                "sentinelCount": int(points),
                "sheet": sheet.title,
            })

    if not candidates:
        raise ValueError("Excelから合計行・定点数行・最新週を判定できません")

    candidates.sort(key=lambda item: (item["year"], item["week"]))
    latest = candidates[-1]
    print("採用した最新週:", latest)
    return latest


def iso_week_period(year, week):
    monday = date.fromisocalendar(year, week, 1)
    sunday = date.fromisocalendar(year, week, 7)
    return f"{monday.month}月{monday.day}日～{sunday.month}月{sunday.day}日"


def get_influenza_data(previous_status):
    excel_url = find_flu_excel_url(fetch_text(FLU_TOP_URL))
    print("インフルエンザExcel:", excel_url)
    workbook = load_workbook(
        io.BytesIO(fetch_bytes(excel_url)),
        data_only=True,
        read_only=True,
    )
    try:
        latest = find_latest_record(workbook)
    finally:
        workbook.close()

    previous_value = previous_status.get("influenza", {}).get("perSentinel")
    difference = None
    if previous_value is not None:
        try:
            difference = round(latest["value"] - float(previous_value), 2)
        except (TypeError, ValueError):
            pass

    return {
        "level": determine_flu_level(latest["value"]),
        "week": f"第{latest['week']}週",
        "period": iso_week_period(latest["year"], latest["week"]),
        "perSentinel": latest["value"],
        "previousPerSentinel": previous_value,
        "difference": difference,
        "reportedCases": latest["reportedCases"],
        "sentinelCount": latest["sentinelCount"],
        "trend": "東京都公式週報",
        "provisional": True,
        "dataStatus": "取得成功",
        "sourceUrl": excel_url,
    }


def main():
    previous = load_previous_status()
    previous_weather = previous.get("weather", {})
    errors = []
    now_text = datetime.now(JST).strftime("%Y-%m-%d %H:%M")

    warning_debug = None
    try:
        warning_debug = get_vpws50_weather_alerts()
        wind_active = warning_debug["windActive"]
        dry_active = warning_debug["dryActive"]
    except Exception as error:
        print("VPWS50取得エラー:", repr(error))
        errors.append("警報注意報")
        wind_active = previous_weather.get("wind", {}).get("active", False)
        dry_active = previous_weather.get("dry", {}).get("active", False)
        warning_debug = {
            "sourceType": VPWS50_TYPE,
            "error": type(error).__name__,
            "fallback": "previousStatus",
        }

    try:
        minimum_temperature, minimum_date, temperature_source = (
            get_overnight_minimum_temperature()
        )
        cold_active = minimum_temperature is not None and minimum_temperature <= 3
        cold_note = (
            f"今夜～明朝 {minimum_temperature:g}℃"
            if minimum_temperature is not None
            else "今夜～明朝の最低気温を確認中"
        )
        cold_debug = {
            "targetDate": minimum_date.isoformat(),
            "minimumTemperature": minimum_temperature,
            "source": temperature_source,
        }
    except Exception as error:
        print("気温取得エラー:", repr(error))
        errors.append("気温予報")
        cold_active = previous_weather.get("cold", {}).get("active", False)
        cold_note = "今夜～明朝の最低気温を確認中"
        cold_debug = {"error": type(error).__name__, "fallback": "previousStatus"}

    try:
        influenza = get_influenza_data(previous)
    except Exception as error:
        print("インフルエンザ取得エラー:", repr(error))
        errors.append("感染症情報")
        old = previous.get("influenza", {})
        if old.get("perSentinel") is not None:
            influenza = {**old, "dataStatus": "取得失敗・前回値を維持"}
        else:
            influenza = {
                "level": "確認中",
                "week": "最新発表",
                "period": "東京都",
                "perSentinel": None,
                "previousPerSentinel": None,
                "difference": None,
                "reportedCases": None,
                "sentinelCount": None,
                "trend": "公式情報を確認中",
                "provisional": True,
                "dataStatus": "取得失敗",
                "sourceUrl": FLU_TOP_URL,
            }

    status = {
        "updated": now_text,
        "sourceStatus": "正常" if not errors else "一部取得失敗：" + "・".join(errors),
        "influenza": influenza,
        "weather": {
            "wind": {
                "active": wind_active,
                "normalText": "情報なし",
                "alertText": "強風注意",
                "note": "飛散・揚重確認",
            },
            "dry": {
                "active": dry_active,
                "normalText": "情報なし",
                "alertText": "火気注意",
                "note": "消火確認を徹底",
            },
            "cold": {
                "active": cold_active,
                "normalText": "情報なし",
                "alertText": "凍結注意",
                "note": cold_note,
            },
        },
        "debug": {
            "warning": warning_debug,
            "cold": cold_debug,
        },
    }

    STATUS_FILE.parent.mkdir(parents=True, exist_ok=True)
    temporary_file = STATUS_FILE.with_suffix(".tmp")
    temporary_file.write_text(
        json.dumps(status, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    json.loads(temporary_file.read_text(encoding="utf-8"))
    temporary_file.replace(STATUS_FILE)
    print(json.dumps(status, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
